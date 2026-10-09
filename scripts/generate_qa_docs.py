import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_BREAK

def create_simple_docx():
    """Simple.docx (Standard tables)."""
    doc = Document()
    doc.add_heading('Simple Document with Standard Tables', 0)

    doc.add_heading('Table 1: Basic Data', level=1)
    table = doc.add_table(rows=4, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'ID'
    hdr_cells[1].text = 'Name'
    hdr_cells[2].text = 'Role'

    data = [
        ('1', 'Alice', 'Engineer'),
        ('2', 'Bob', 'Manager'),
        ('3', 'Charlie', 'Designer')
    ]
    for i, row in enumerate(data):
        cells = table.rows[i+1].cells
        for j, val in enumerate(row):
            cells[j].text = val

    doc.add_paragraph("\nSome text between tables.\n")

    doc.add_heading('Table 2: Financials', level=1)
    table2 = doc.add_table(rows=3, cols=2)
    table2.rows[0].cells[0].text = 'Quarter'
    table2.rows[0].cells[1].text = 'Revenue'
    table2.rows[1].cells[0].text = 'Q1'
    table2.rows[1].cells[1].text = '$10,000'
    table2.rows[2].cells[0].text = 'Q2'
    table2.rows[2].cells[1].text = '$12,000'

    doc.save('Simple.docx')
    print("Created Simple.docx")

def create_complex_docx():
    """Complex.docx (Merged cells, vertical splits)."""
    doc = Document()
    doc.add_heading('Complex Document', 0)

    # Merged Cells
    doc.add_heading('Table 1: Merged Cells', level=1)
    table = doc.add_table(rows=4, cols=4)
    table.style = 'Table Grid'

    # Header A and B merged into one
    row0 = table.rows[0]
    row0.cells[0].merge(row0.cells[1])
    row0.cells[0].text = 'Header A & B'
    row0.cells[2].text = 'Header C'
    row0.cells[3].text = 'Header D'

    # Vertical Merge
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 0).text = 'Vertical Merge'

    table.cell(1, 1).text = 'Data 1'
    table.cell(2, 1).text = 'Data 2'

    # Fill rest
    for r in range(1, 4):
        for c in range(2, 4):
            table.cell(r, c).text = f'R{r}C{c}'

    doc.add_paragraph("\n")

    # Vertical Splits / Complex Header
    doc.add_heading('Table 2: Complex Header', level=1)
    table2 = doc.add_table(rows=5, cols=3)
    table2.style = 'Table Grid'

    # Row 1: "Main Category" spans all
    table2.cell(0, 0).merge(table2.cell(0, 2))
    table2.cell(0, 0).text = "Main Category"

    # Row 2: Subcats
    table2.cell(1, 0).text = "Sub 1"
    table2.cell(1, 1).text = "Sub 2"
    table2.cell(1, 2).text = "Sub 3"

    # Data
    for r in range(2, 5):
        for c in range(3):
            table2.cell(r, c).text = f"Val {r}-{c}"

    doc.save('Complex.docx')
    print("Created Complex.docx")

def create_empty_docx():
    """Empty.docx (No tables)."""
    doc = Document()
    doc.add_heading('Empty Document', 0)
    doc.add_paragraph("This document contains text but no tables.")
    doc.add_paragraph("Just some paragraphs here and there.")
    doc.save('Empty.docx')
    print("Created Empty.docx")

def create_huge_docx():
    """Huge.docx (50+ pages)."""
    doc = Document()
    doc.add_heading('Huge Document', 0)

    # Add 55 pages with 1 table each
    for i in range(55):
        doc.add_heading(f'Page {i+1}', level=1)
        doc.add_paragraph(f"Content for page {i+1}...")

        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = f"P{i+1}-Col1"
        table.rows[0].cells[1].text = f"P{i+1}-Col2"
        table.rows[1].cells[0].text = "Data"
        table.rows[1].cells[1].text = "Data"

        doc.add_paragraph("\n")
        # Force page break
        if i < 54:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.save('Huge.docx')
    print("Created Huge.docx")

if __name__ == "__main__":
    create_simple_docx()
    create_complex_docx()
    create_empty_docx()
    create_huge_docx()
